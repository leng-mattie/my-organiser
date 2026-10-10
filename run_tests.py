"""Automated run of TESTING.md with headless Chromium (Playwright) + openpyxl.
Usage:  python3 tests/run_tests.py   (from the my-organiser folder's parent or anywhere)
Serves the app from a sub-folder over http to mimic GitHub Pages."""
import datetime as dt, http.server, json, os, re, shutil, sys, tempfile, threading, functools
from playwright.sync_api import sync_playwright

APP = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ROOT = tempfile.mkdtemp()
os.symlink(APP, os.path.join(ROOT, 'my-organiser'))
OUT = tempfile.mkdtemp()
Handler = functools.partial(http.server.SimpleHTTPRequestHandler, directory=ROOT)
Handler.log_message = lambda *a: None
srv = http.server.ThreadingHTTPServer(('127.0.0.1', 0), Handler)
threading.Thread(target=srv.serve_forever, daemon=True).start()
URL = f'http://127.0.0.1:{srv.server_address[1]}/my-organiser/'

TODAY = dt.date(2026, 10, 12)          # a Monday; clock is fixed at 10:00 London time
def d(off): return (TODAY + dt.timedelta(days=off)).isoformat()
results = []
def check(n, name, ok, detail=''):
    results.append((n, name, bool(ok), detail))
    print(('PASS' if ok else 'FAIL'), f'#{n}', name, ('— ' + detail) if detail and not ok else '')

def new_page(browser, width, height, dark=False, block_storage=False):
    ctx = browser.new_context(viewport={'width': width, 'height': height}, accept_downloads=True,
                              color_scheme='dark' if dark else 'light', timezone_id='Europe/London', locale='en-GB',
                              service_workers='allow')
    page = ctx.new_page()
    page.clock.install(time=dt.datetime(2026, 10, 12, 9, 0, tzinfo=dt.timezone.utc))
    errors = []
    page.on('console', lambda m: errors.append(m.text) if m.type == 'error' else None)
    page.on('pageerror', lambda e: errors.append(str(e)))
    page.on('dialog', lambda dlg: dlg.accept())
    if block_storage:
        page.add_init_script("Object.defineProperty(window,'localStorage',{get(){throw new DOMException('blocked','SecurityError')}})")
    page.goto(URL)
    return ctx, page, errors

def add_task(page, title, off, prio='Medium', time=''):
    page.goto(URL + '#add')
    page.fill('#f-title', title)
    page.select_option('#f-dest', 'timetable')
    if off is None: page.select_option('#f-hasdl', 'No')
    else:
        page.select_option('#f-hasdl', 'Yes'); page.fill('#f-date', d(off))
        if time: page.fill('#f-time', time)
    page.select_option('#f-priority', prio)
    page.click('#f-submit')

def add_budget(page, title, typ, amount, account='HSBC', cat='Food', date=None, frm=None, to=None, due=None):
    page.goto(URL + '#add')
    page.fill('#f-title', title)
    page.select_option('#f-dest', 'budget')
    page.select_option('#f-type', typ)
    if typ == 'transfer':
        page.select_option('#f-from', frm); page.select_option('#f-to', to)
    else:
        page.select_option('#f-account', account); page.select_option('#f-category', cat)
    page.fill('#f-amount', amount)
    if date: page.fill('#f-bdate', date)
    if due: page.fill('#f-due', due)
    page.click('#f-submit')

def no_hscroll(page):
    return page.evaluate('document.documentElement.scrollWidth <= window.innerWidth')

with sync_playwright() as p:
    browser = p.chromium.launch()

    for (W, H, label) in [(375, 812, 'phone'), (1280, 800, 'laptop')]:
        print(f'\n=== {label} {W}x{H} ===')
        ctx, page, errors = new_page(browser, W, H)
        tag = f' [{label}]'

        # 1 empty state
        ok = page.is_visible('#page-add') and page.locator('#f-title').is_visible()
        page.goto(URL + '#timetable')
        txt = page.inner_text('#tt-content')
        check(1, 'Empty app opens calmly' + tag, ok and 'clean slate' in txt and 'Nothing to do yet' in txt and no_hscroll(page))

        # 2 empty title
        page.goto(URL + '#add'); page.click('#f-submit')
        check(2, 'Empty title is rejected kindly' + tag,
              'Please type what you need to add' in page.inner_text('#e-title') and page.evaluate('document.activeElement.id') == 'f-title')

        # 3 deadline yes but no date; then No hides date
        page.fill('#f-title', 'Read chapter 4'); page.select_option('#f-hasdl', 'Yes'); page.click('#f-submit')
        e3 = page.inner_text('#e-date')
        page.select_option('#f-hasdl', 'No'); hidden = not page.is_visible('#f-date')
        page.click('#f-submit')
        check(3, 'Missing date caught; No-deadline hides date and saves' + tag,
              'Pick a deadline date' in e3 and hidden and 'Added' in page.inner_text('#add-confirm'))

        # 4 add due today + confirmation link
        add_task(page, 'Submit HSBC assessment', 0, 'High', '14:00')
        conf = page.inner_text('#add-confirm')
        page.click('#add-confirm a')
        page.wait_for_selector('#tt-panel .item')      # hash navigation is asynchronous
        tt = page.inner_text('#tt-panel')
        check(4, 'Due-today item lands in Today with link' + tag,
              'Timetable' in conf and page.url.endswith('#timetable') and 'Submit HSBC assessment' in tt and 'Due today' in tt)

        # 5 buckets and lines
        add_task(page, 'Essay plan', 1, 'Medium'); add_task(page, 'Shadowing day', 3, 'Low', '09:00')
        add_task(page, 'Dissertation', 10, 'High'); add_task(page, 'Library fine', -2, 'Low')
        page.goto(URL + '#timetable')
        today_txt = page.inner_text('#tt-panel')
        first_today = page.locator('#tt-panel .item .title').first.inner_text()
        page.click('#tab-week'); week = page.inner_text('#tt-panel')
        page.click('#tab-later'); later = page.inner_text('#tt-panel')
        page.click('#tab-nodl'); nodl = page.inner_text('#tt-panel')
        check(5, 'Day lines + tabs + overdue first' + tag,
              '2 days overdue' in today_txt and first_today == 'Library fine' and 'Due tomorrow' in week and '3 days left' in week
              and '10 days left' in later and 'No deadline' in nodl and 'Read chapter 4' in nodl,
              f'first={first_today}')

        # 6 top 3
        top = page.inner_text('.top3')
        check(6, 'Top 3 with reasons' + tag, page.locator('.top3 li').count() == 3 and 'overdue' in top and 'Due today' in top)

        # 7 priority icon + text
        page.click('#tab-today')
        prio = page.locator('#tt-panel .prio').first
        check(7, 'Priority shown with icon AND text' + tag, prio.locator('svg').count() == 1 and 'priority' in prio.inner_text())

        # 8 tick done + undo
        page.click('#tt-panel .tick >> nth=0')
        page.click('#tab-done'); in_done = 'Library fine' in page.inner_text('#tt-panel')
        page.click('#toast-undo'); page.click('#tab-done'); undone = 'Library fine' not in page.inner_text('#tt-panel')
        check(8, 'Tick done and Undo' + tag, in_done and undone)

        # 9 edit
        page.click('#tab-later')
        page.click('#tt-panel details.more summary'); page.click('#tt-panel [data-action=edit-task]')
        same_form = page.url.endswith('#add') and page.input_value('#f-title') == 'Dissertation' and page.input_value('#f-priority') == 'High'
        page.select_option('#f-priority', 'Low'); page.click('#f-submit')
        page.click('#tab-later')
        edited = 'Low priority' in page.inner_text('#tt-panel')
        page.click('#tt-panel details.more summary'); page.click('#tt-panel [data-action=edit-task]')
        page.fill('#f-title', 'SHOULD NOT SAVE'); page.click('#f-cancel')
        check(9, 'Edit uses same drop-downs; cancel changes nothing' + tag,
              same_form and edited and 'SHOULD NOT SAVE' not in page.inner_text('#tt-content') and page.url.endswith('#timetable'))

        # 10 delete + undo
        page.click('#tab-later')
        page.click('#tt-panel details.more summary'); page.click('#tt-panel [data-action=delete-task]')
        gone = 'Dissertation' not in page.inner_text('#tt-panel')
        page.click('#toast-undo'); back = 'Dissertation' in page.inner_text('#tt-panel')
        check(10, 'Delete and Undo' + tag, gone and back)

        # 11 overload warning
        before = page.locator('.card.attn').count()
        for i in range(6): add_task(page, f'High thing {i}', 1 + (i % 3), 'High')
        page.goto(URL + '#timetable')
        check(11, 'Overload warning at >5 high items in 3 days' + tag,
              before == 0 and 'That’s a lot at once' in page.inner_text('#tt-content'))

        # 12 ICS
        with page.expect_download() as dl: page.click('[data-action=ics-all]')
        path = os.path.join(OUT, f'all-{label}.ics'); dl.value.save_as(path)
        ics = open(path, 'rb').read().decode('utf-8')
        page.click('#tab-today'); page.click('#tt-panel details.more summary >> nth=0')
        with page.expect_download() as dl1: page.click('#tt-panel [data-action=ics-one] >> nth=0')
        one = open(dl1.value.path(), 'rb').read().decode()
        lines = ics.split('\r\n')
        ok12 = (ics.startswith('BEGIN:VCALENDAR\r\n') and ics.endswith('END:VCALENDAR\r\n') and '\n' not in ics.replace('\r\n', '')
                and f"DTSTART;VALUE=DATE:{d(1).replace('-', '')}" in ics
                and f"DTSTART;TZID=Europe/London:{d(0).replace('-', '')}T140000" in ics and 'BEGIN:VTIMEZONE' in ics
                and all(len(l.encode()) <= 75 for l in lines) and ics.count('BEGIN:VEVENT') == ics.count('END:VEVENT') == 11
                and one.count('BEGIN:VEVENT') == 1 and 'one-time export' in page.inner_text('#tt-content'))
        check(12, 'Apple Calendar .ics (all-day + timed, folded CRLF)' + tag, ok12, f"events={ics.count('BEGIN:VEVENT')}")

        # 13 money validation
        msgs = []
        for amt in ['-5', '0', 'abc', '12.345', '2000000', '']:
            add_budget(page, 'Lunch', 'out', amt)
            msgs.append(page.inner_text('#e-amount'))
        check(13, 'Bad amounts rejected with clear messages' + tag,
              'negative' in msgs[0] and 'more than £0' in msgs[1] and 'Numbers only' in msgs[2] and 'two decimal' in msgs[3]
              and '1,000,000' in msgs[4] and 'enter an amount' in msgs[5], ' | '.join(msgs))

        # 14 balances and transfer
        page.goto(URL + '#settings')
        page.fill('#s-starthsbc', '500'); page.fill('#s-startmonzo', '1,000.50'); page.fill('#s-budget', '400'); page.fill('#s-goal', '2000')
        page.click('#money-form button[type=submit]')
        add_budget(page, 'Part-time job', 'in', '250', 'HSBC', 'Other')
        add_budget(page, 'Groceries', 'out', '45.20', 'HSBC', 'Food')
        add_budget(page, 'Coffee', 'out', '3.10', 'HSBC', 'Food', date=d(-1))
        add_budget(page, 'To savings', 'transfer', '100', frm='HSBC', to='Monzo')
        page.goto(URL + '#budget')
        bud = page.inner_text('#bud-content')
        hsbc = 500 + 250 - 45.20 - 3.10 - 100
        monzo = 1000.50 + 100
        chart = page.inner_text('.chart')
        check(14, 'Balances add up; transfer not spending' + tag,
              f'£{hsbc:,.2f}' in bud and f'£{monzo:,.2f}' in bud and f'£{hsbc + monzo:,.2f}' in bud
              and 'Food' in chart and '£48.30' in chart and 'Transfer HSBC → Monzo' in bud and '£100' not in chart)

        # 15 from = to
        page.goto(URL + '#add'); page.select_option('#f-dest', 'budget'); page.select_option('#f-type', 'transfer')
        page.select_option('#f-from', 'Monzo'); flipped = page.input_value('#f-to') == 'HSBC'
        page.evaluate("document.querySelector('#f-to').value='Monzo'")
        page.fill('#f-title', 'X'); page.fill('#f-amount', '5'); page.click('#f-submit')
        check(15, 'Transfer From/To must differ' + tag, flipped and 'must be different' in page.inner_text('#e-to'))

        # 16 budget bar, goal, upcoming, mark paid
        add_budget(page, 'Rent', 'upcoming', '300', 'HSBC', 'Bills', due=d(5))
        page.goto(URL + '#budget')
        bud = page.inner_text('#bud-content')
        after = hsbc - 300
        stats = page.inner_text('.stats')
        bars = page.locator('.progress').count()
        ok16a = (f'£{after:,.2f}' in bud and '5 days left' in bud and bars == 2 and '£45.20' in stats and '£48.30' in stats)
        page.click('#h-up >> xpath=.. >> details.more summary'); page.click('[data-action=pay-up]')
        bud2 = page.inner_text('#bud-content')
        check(16, 'Budget bar, goal, upcoming + after-upcoming, mark paid' + tag,
              ok16a and 'Nothing coming up' in bud2 and f'£{after:,.2f}' in page.inner_text('#bud-content'), stats)

        # 17 Excel export
        page.goto(URL + '#settings')
        with page.expect_download() as dlx: page.click('[data-action=export-xlsx]')
        xpath = os.path.join(OUT, f'export-{label}.xlsx'); dlx.value.save_as(xpath)
        import openpyxl
        wb = openpyxl.load_workbook(xpath)
        dv = {ws.title: {str(v.sqref): v.formula1 for v in ws.data_validations.dataValidation} for ws in wb}
        tts = wb['Timetable']
        hdr = [c.value for c in tts[1]]
        txs = wb['Budget transactions']
        amounts = [r[6] for r in txs.iter_rows(min_row=2, values_only=True)]
        ok17 = (wb.sheetnames == ['Timetable', 'Budget transactions', 'Upcoming expenses', 'Settings']
                and hdr == ['Title', 'Has deadline', 'Deadline', 'Time', 'Priority', 'Note', 'Done']
                and dv['Timetable'].get('E2:E1000') == '"High,Medium,Low"' and dv['Timetable'].get('B2:B1000') == '"Yes,No"'
                and dv['Budget transactions'].get('C2:C1000') == '"Money out,Money in,Transfer"'
                and dv['Budget transactions'].get('D2:D1000') == '"HSBC,Monzo"'
                and dv['Budget transactions'].get('H2:H1000', '').startswith('Settings!$D$2')
                and 300 in amounts and 45.2 in amounts
                and tts['C2'].number_format == 'dd/mm/yyyy' and isinstance(tts['C2'].value, dt.datetime)
                and wb['Settings']['B2'].value == 500)
        check(17, 'Excel export: 4 sheets, data, drop-downs, UK formats' + tag, ok17, json.dumps(dv)[:300])

        # 18 import back: edit in openpyxl (simulates Excel), add good + bad row; also CSV
        tts.append(['New from Excel', 'Yes', dt.datetime(2026, 10, 20), '', 'High', '', 'No'])
        tts.append(['Bad row', 'Yes', dt.datetime(2026, 10, 20), '', 'Urgent', '', 'No'])
        tts.append(['', 'Yes', 'not a date', '', 'Low', '', 'No'])
        txs.append([dt.datetime(2026, 10, 11), 'Bus pass', 'Money out', 'HSBC', None, None, -5, 'Travel'])
        wb.save(xpath)
        page.select_option('#s-importmode', 'add')
        page.set_input_files('#import-file', xpath)
        page.wait_for_selector('#import-report .confirm, #import-report .err-summary')
        rep = page.inner_text('#import-report')
        csvp = os.path.join(OUT, 'tt.csv')
        open(csvp, 'w', encoding='utf-8-sig').write('Title,Has deadline,Deadline,Time,Priority,Note,Done\r\n"Call mum, Sunday",Yes,18/10/2026,,Low,"multi\nline",No\r\nNo prio,No,,,,,No\r\n')
        page.evaluate("document.querySelector('#import-report').innerHTML = ''")
        page.set_input_files('#import-file', csvp)
        page.wait_for_selector('#import-report .confirm, #import-report .err-summary')
        rep2 = page.inner_text('#import-report')
        page.goto(URL + '#timetable'); page.click('#tab-week')
        check(18, 'Excel + CSV import: adds good rows, skips dups, lists rejects' + tag,
              'Added 1 timetable item' in rep and 'skipped because they were already' in rep
              and re.search(r'Timetable, row \d+: Priority', rep) and 'Title/Description' in rep and 'Deadline' in rep
              and 'Amount — Amounts can’t be negative' in rep and '1 row rejected' not in rep
              and 'Added 1 timetable item' in rep2 and 'Priority' in rep2 and 'Call mum, Sunday' in page.inner_text('#tt-panel'),
              rep.replace('\n', ' | ')[:500])

        # 19 JSON backup round trip + bad file
        page.goto(URL + '#settings')
        before_state = page.evaluate("localStorage.getItem('myOrganiser.v1')")
        with page.expect_download() as dlj: page.click('[data-action=export-json]')
        jpath = os.path.join(OUT, 'backup.json'); dlj.value.save_as(jpath)
        page.click('[data-action=clear-all]')
        cleared = '0 timetable items' in page.inner_text('#storage-status')
        page.set_input_files('#import-json', jpath)
        page.wait_for_selector('#json-report .confirm')
        after_state = json.loads(page.evaluate("localStorage.getItem('myOrganiser.v1')"))
        b = json.loads(before_state)
        same = (len(after_state['tasks']) == len(b['tasks']) and after_state['transactions'] == b['transactions']
                and after_state['settings']['startMonzo'] == 100050)
        badp = os.path.join(OUT, 'bad.json'); open(badp, 'w').write('{"hello": 1}')
        page.set_input_files('#import-json', badp)
        page.wait_for_selector('#json-report .err-summary')
        check(19, 'Backup export / clear / restore; bad file refused' + tag,
              cleared and same and 'doesn’t look like a My Organiser backup' in page.inner_text('#json-report'))

        # 20 a) no horizontal scroll on every page
        hs = []
        for pg in ['add', 'timetable', 'budget', 'settings']:
            page.goto(URL + '#' + pg); hs.append(no_hscroll(page))
        page.goto(URL + '#add'); page.select_option('#f-dest', 'budget'); page.select_option('#f-type', 'transfer'); hs.append(no_hscroll(page))
        # b) font size >= 16 for visible text
        small = page.evaluate("""() => { const bad=[]; document.querySelectorAll('main *, nav *').forEach(el=>{
            if(!el.offsetParent || !el.childNodes.length) return;
            const own=[...el.childNodes].some(n=>n.nodeType===3&&n.textContent.trim()); if(!own) return;
            if (el.closest('svg')) return;
            const fs=parseFloat(getComputedStyle(el).fontSize); if(fs<16 && !el.closest('.nav')) bad.push(el.tagName+':'+fs+':'+el.textContent.trim().slice(0,20)); }); return bad; }""")
        # c) tap targets >= 44px
        tiny = page.evaluate("""() => [...document.querySelectorAll('button, .nav a, a.btn, select, input, summary')].filter(e=>e.offsetParent && !e.closest('.vh'))
              .map(e=>{const r=e.getBoundingClientRect(); return [e.id||e.textContent.trim().slice(0,15), Math.round(r.height)]}).filter(x=>x[1]<44)""")
        check('20a', 'No sideways scroll; text ≥16px; tap targets ≥44px' + tag, all(hs) and not small and not tiny, f'{hs} {small[:5]} {tiny[:8]}')

        # 20 b) theme toggle
        page.goto(URL + '#timetable')
        bg1 = page.evaluate("getComputedStyle(document.body).backgroundColor")
        page.click('#theme-btn'); bg2 = page.evaluate("getComputedStyle(document.body).backgroundColor")
        page.click('#theme-btn'); bg3 = page.evaluate("getComputedStyle(document.body).backgroundColor")
        check('20b', 'Manual dark/light toggle' + tag, bg1 == bg3 and bg1 != bg2, f'{bg1} {bg2}')

        # 20 c) midnight rollover with the app left open
        page.clock.set_system_time(dt.datetime(2026, 10, 12, 22, 59, 30, tzinfo=dt.timezone.utc))   # 23:59:30 BST
        page.goto(URL + '#timetable'); page.click('#tab-week')
        pre = 'Due tomorrow' in page.inner_text('#tt-panel')
        page.clock.run_for(90 * 1000)    # passes midnight London time; app checks every 30 s
        page.click('#tab-today')
        post = page.inner_text('#tt-panel')
        check('20c', 'Midnight rollover updates “Due tomorrow” → “Due today”' + tag, pre and 'Essay plan' in post and 'Due today' in post)

        check('20d', 'No console errors' + tag, not errors, ' | '.join(errors[:5]))
        page.screenshot(path=os.path.join(OUT, f'{label}-timetable.png'), full_page=True)
        page.goto(URL + '#budget'); page.screenshot(path=os.path.join(OUT, f'{label}-budget.png'), full_page=True)
        page.goto(URL + '#add'); page.screenshot(path=os.path.join(OUT, f'{label}-add.png'), full_page=True)
        page.goto(URL + '#settings'); page.screenshot(path=os.path.join(OUT, f'{label}-settings.png'), full_page=True)
        ctx.close()

    # Dark mode follows the system
    ctx, page, errors = new_page(browser, 375, 812, dark=True)
    add_task(page, 'Dark check', 0, 'High'); page.goto(URL + '#timetable')
    bg = page.evaluate("getComputedStyle(document.body).backgroundColor")
    page.screenshot(path=os.path.join(OUT, 'phone-dark.png'), full_page=True)
    check('20e', 'Dark mode follows the system', bg == 'rgb(18, 22, 20)', bg)

    # Offline: service worker serves everything
    page.evaluate("navigator.serviceWorker.ready.then(() => true)")
    page.reload()
    page.evaluate("new Promise(r => { if (navigator.serviceWorker.controller) r(true); else navigator.serviceWorker.addEventListener('controllerchange', () => r(true)); })")
    ctx.set_offline(True)
    page.goto(URL)
    offline_ok = page.is_visible('#f-title') and page.evaluate("typeof JSZip") == 'function'
    page.goto(URL + 'index.html#timetable'); offline_ok = offline_ok and 'Dark check' in page.inner_text('#tt-content')
    cached = page.evaluate("caches.keys().then(k=>caches.open(k[0])).then(c=>c.keys()).then(r=>r.map(x=>x.url))")
    manifest = page.evaluate("fetch('manifest.webmanifest').then(r=>r.json())")
    check('20f', 'Works offline from a sub-folder (service worker)', offline_ok and len(cached) >= 7, str(cached))
    check('20g', 'Manifest correct', manifest['name'] == 'My Organiser' and manifest['display'] == 'standalone'
          and manifest['start_url'] == './' and manifest['scope'] == './')
    ctx.close()

    # Storage blocked → in-memory fallback + banner
    ctx, page, errors = new_page(browser, 375, 812, block_storage=True)
    add_task(page, 'Memory only', None, 'Low'); page.goto(URL + '#timetable'); page.click('#tab-nodl')
    check('20h', 'Blocked storage → works in memory with a warning',
          'Memory only' in page.inner_text('#tt-panel') and page.is_visible('#storage-banner') and not errors, ' | '.join(errors))
    ctx.close()

    # Opening by double-click (file://)
    ctx = browser.new_context(viewport={'width': 375, 'height': 812}, accept_downloads=True)
    page = ctx.new_page(); ferr = []
    page.on('pageerror', lambda e: ferr.append(str(e)))
    page.on('console', lambda m: ferr.append(m.text) if m.type == 'error' and 'manifest' not in m.text.lower() else None)
    page.on('dialog', lambda dlg: dlg.accept())
    page.goto('file://' + os.path.join(APP, 'index.html'))
    add_task_url = 'file://' + os.path.join(APP, 'index.html')
    page.fill('#f-title', 'File mode'); page.select_option('#f-hasdl', 'No'); page.click('#f-submit')
    page.goto(add_task_url + '#settings')
    with page.expect_download() as dlf: page.click('[data-action=export-xlsx]')
    check('20i', 'Works when double-clicked (file://), incl. Excel export',
          'File mode' in page.evaluate("localStorage.getItem('myOrganiser.v1')") and dlf.value.suggested_filename.endswith('.xlsx') and not ferr, ' | '.join(ferr))
    ctx.close()
    browser.close()

print('\nSUMMARY:', sum(r[2] for r in results), '/', len(results), 'passed. Output files in', OUT)
fails = [r for r in results if not r[2]]
for f in fails: print('FAILED', f)
sys.exit(1 if fails else 0)
